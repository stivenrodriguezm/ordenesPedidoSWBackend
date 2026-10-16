from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from openpyxl import load_workbook
from rest_framework.test import APIClient
from ordenes.models import ReciboCaja, PagoReciboCaja, ComprobanteEgreso, RegistroAuditoria
from ordenes.tests.utils import LottusTestCase
from suministros.models import Inventario, GrupoInventario, Categoria, Subcategoria, Sede, Zona


class ReportesOperativosTests(LottusTestCase):
    def setUp(self):
        self.admin = self.crear_usuario('admin-reportes', role='administrador')
        self.usuario = self.crear_usuario('lector-reportes')
        self.otro = self.crear_usuario('otro-reportes')
        self.permisos('vendedor', ['VER_PROPIAS_VENTAS', 'VER_RECIBOS', 'VER_INVENTARIO', 'VER_COMPROBANTES_EGRESO'])
        self.c = self.api(self.admin)
        self.limitado = self.api(self.usuario)
        self.proveedor_a = self.proveedor('=PROVEEDOR ficticio')
        self.ref = self.referencia(self.proveedor_a, 'Mesa ficticia')
        self.v = self.venta(self.usuario, valor='1000.25', fecha=date(2026, 10, 6), sede='CENTRO')
        self.ajena = self.venta(self.otro, fecha=date(2026, 9, 5), sede='NORTE')
        self.rc = ReciboCaja.objects.create(id=9001, venta=self.v, nota='=SUM(1,1)')
        PagoReciboCaja.objects.create(recibo=self.rc, metodo_pago='Efectivo', valor='100.25', estado='Confirmado')
        self.ce = ComprobanteEgreso.objects.create(id=9101, proveedor=self.proveedor_a, valor='50.25', medio_pago='Transferencia', concepto='=prueba', estado='Por Confirmar Pago')
        self.i = Inventario.objects.create(id_referencia='001A', referencia=self.ref, costo_especifico='765.25', observacion='=HYPERLINK("x")')

    def exportar(self, tipo, params=None, cliente=None):
        from urllib.parse import urlencode
        return (cliente or self.c).post(f'/api/reportes-operativos/{tipo}/exportar/?' + urlencode(params or {}))

    def filas(self, response):
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        return list(load_workbook(BytesIO(response.content), data_only=False).active.values)

    def test_api_exige_sesion_y_permiso_incluso_sin_boton(self):
        for tipo in ('ventas', 'recibos', 'egresos', 'inventario'):
            self.assertEqual(self.exportar(tipo, cliente=APIClient()).status_code, 401)
        for tipo in ('ventas', 'recibos', 'egresos'):
            self.assertEqual(self.exportar(tipo, cliente=self.limitado).status_code, 403)
        self.permisos('vendedor', [])
        self.assertEqual(self.exportar('inventario', cliente=self.limitado).status_code, 403)
        self.assertEqual(self.exportar('desconocido').status_code, 404)
        self.assertFalse(RegistroAuditoria.objects.exists())

    def test_ventas_filtro_periodo_estricto_y_total_completo_sin_division(self):
        self.v.vendedores_compartidos.add(self.otro)
        r = self.exportar('ventas', {'periods': '10-2026', 'page': '2', 'page_size': '1'})
        filas = self.filas(r)
        self.assertEqual(len(filas), 2)
        self.assertEqual(filas[1][0], str(self.v.pk))
        self.assertEqual(Decimal(str(filas[1][-1])), Decimal('1000.25'))
        listado = self.c.get('/api/ventas/?periods=10-2026&page=1&page_size=1')
        self.assertEqual(listado.data['count'], 1)
        self.assertEqual(listado.data['results'][0]['id'], int(filas[1][0]))

    def test_ventas_excluir_vendedores_respeta_compartidas_e_inactivos(self):
        inactivo = self.crear_usuario('inactivo-reportes')
        inactivo.is_active = False
        inactivo.save(update_fields=['is_active'])
        de_inactivo = self.venta(inactivo)
        self.v.vendedores_compartidos.add(self.otro)
        def ids(params):
            r = self.c.get('/api/ventas/', params)
            self.assertEqual(r.status_code, 200, r.data)
            return {v['id'] for v in r.data}
        todas = ids({})
        self.assertIn(de_inactivo.pk, todas)
        # Quitar al vendedor principal no oculta la venta compartida con alguien que sigue marcado.
        self.assertIn(self.v.pk, ids({'excluir_vendedores': str(self.usuario.pk)}))
        self.assertNotIn(self.v.pk, ids({'excluir_vendedores': f'{self.usuario.pk},{self.otro.pk}'}))
        self.assertNotIn(self.ajena.pk, ids({'excluir_vendedores': str(self.otro.pk)}))
        self.assertEqual(ids({'excluir_vendedores': str(inactivo.pk)}), todas - {de_inactivo.pk})
        self.assertEqual(self.c.get('/api/ventas/', {'excluir_vendedores': 'x'}).status_code, 400)
        informe = {v['id']: v for v in self.c.get('/api/ventas/', {'is_report': 'true'}).data}
        self.assertEqual(informe[self.v.pk]['vendedores_compartidos_nombres'], [self.otro.first_name])

    def test_ventas_exporta_mas_de_una_pagina_y_busqueda_prioriza_como_listado(self):
        for _ in range(32):
            self.venta(self.usuario)
        filas = self.filas(self.exportar('ventas', {'page': '2', 'page_size': '1'}))
        self.assertEqual(len(filas), 35)
        r = self.exportar('ventas', {'search': str(self.ajena.pk), 'estado': 'imposible'})
        self.assertEqual(self.filas(r)[1][0], str(self.ajena.pk))

    def test_listado_y_total_recibos_excluyen_ajenos_incluyen_compartidos(self):
        otro = ReciboCaja.objects.create(id=9002, venta=self.ajena)
        PagoReciboCaja.objects.create(recibo=otro, metodo_pago='Efectivo', valor='700')
        url = '/api/recibos-caja/'
        r = self.limitado.get(url)
        self.assertEqual([x['id'] for x in r.data['results']], [self.rc.pk])
        self.assertEqual(r.data['totales_filtrados']['valor'], Decimal('100.25'))
        self.ajena.vendedores_compartidos.add(self.usuario)
        r = self.limitado.get(url)
        self.assertEqual(r.data['count'], 2)
        self.assertEqual(r.data['totales_filtrados']['valor'], Decimal('800.25'))

    def test_recibos_metodo_filtra_documentos_completos_sin_inventar_estado(self):
        PagoReciboCaja.objects.create(recibo=self.rc, metodo_pago='Davivienda', valor='20', estado='Pendiente')
        filas = self.filas(self.exportar('recibos', {'medio_pago': 'Efectivo'}))
        self.assertEqual(filas[1][4], 'Múltiple')
        self.assertEqual(Decimal(str(filas[1][5])), Decimal('120.25'))
        self.assertEqual(filas[1][-1], 'Parcial')
        self.assertEqual(filas[1][-2], '=SUM(1,1)')

    def test_egresos_permiso_y_filtros_no_confunden_pdf_con_estado(self):
        self.permisos('vendedor', ['VER_COMPROBANTES_EGRESO', 'EXPORTAR_EGRESOS'])
        filas = self.filas(self.exportar('egresos', {'proveedores': self.proveedor_a.pk, 'estados': 'Por Confirmar'}, self.limitado))
        self.assertEqual(len(filas), 2)
        self.assertEqual(filas[1][4], 'Por Confirmar Pago')
        self.assertEqual(Decimal(str(filas[1][5])), Decimal('50.25'))
        self.assertEqual(self.filas(self.exportar('egresos', {'estados': 'Pagado'})), [filas[0]])

    def test_inventario_sin_permiso_no_contiene_columna_costo(self):
        filas = self.filas(self.exportar('inventario', cliente=self.limitado))
        self.assertNotIn('Costo', filas[0])
        self.assertNotIn('765.25', str(filas))
        self.assertEqual(filas[1][0], '001A')
        globales = self.filas(self.exportar('inventario'))
        self.assertEqual(globales[0][-1], 'Costo')
        self.assertEqual(Decimal(str(globales[1][-1])), Decimal('765.25'))

    def test_inventario_filtros_de_grupo_incluyen_componentes_otro_proveedor(self):
        categoria = Categoria.objects.create(nombre='Comedores')
        subcat = Subcategoria.objects.create(nombre='Grupo especial', categoria=categoria)
        grupo = GrupoInventario.objects.create(nombre='Comedor Alpha', categoria=categoria, subcategoria=subcat)
        self.i.grupo = grupo
        self.i.save()
        otro = self.referencia(self.proveedor('Otro'), 'Silla')
        Inventario.objects.create(id_referencia='002B', referencia=otro, grupo=grupo)
        Inventario.objects.create(id_referencia='003C', referencia=otro)
        for filtros in ({'proveedor': self.proveedor_a.pk}, {'categoria': categoria.pk}, {'subcategoria': subcat.pk}, {'search': 'Comedor Alpha'}):
            filas = self.filas(self.exportar('inventario', filtros))
            self.assertEqual([f[0] for f in filas[1:]], ['001A', '002B'])

    def test_inventario_zona_sede_fechas_y_estado(self):
        sede = Sede.objects.create(nombre='Sede ficticia')
        zona = Zona.objects.create(nombre='Zona ficticia', sede=sede)
        self.i.zona = zona
        self.i.fecha_ingreso = date(2026, 10, 6)
        self.i.save()
        self.assertEqual(len(self.filas(self.exportar('inventario', {'sede': sede.pk, 'zona': zona.pk, 'fecha_inicio': '2026-10-01', 'fecha_fin': '2026-10-10', 'disponibilidad': self.i.disponibilidad}))), 2)
        self.assertEqual(len(self.filas(self.exportar('inventario', {'fecha_fin': '2026-09-30'}))), 1)

    def test_archivos_literales_auditoria_sin_pii_y_cabeceras(self):
        for tipo in ('ventas', 'recibos', 'egresos', 'inventario'):
            response = self.exportar(tipo)
            self.assertEqual(response['Cache-Control'], 'no-store')
            self.assertEqual(response['X-Content-Type-Options'], 'nosniff')
            hoja = load_workbook(BytesIO(response.content), data_only=False).active
            for fila in hoja:
                self.assertTrue(all(c.data_type != 'f' for c in fila))
        r = self.exportar('inventario', {'search': 'Mesa ficticia', 'password': 'no-registrar'})
        self.assertEqual(r.status_code, 200)
        detalle = RegistroAuditoria.objects.latest('pk').detalle
        self.assertNotIn('Mesa ficticia', str(detalle))
        self.assertNotIn('password', str(detalle))
        self.assertTrue(detalle['filtros']['search']['huella'])

    def test_limite_sin_truncar_ni_auditar(self):
        with patch('ordenes.reportes_views.LIMITE_EXPORTACION', 1):
            r = self.exportar('ventas')
        self.assertEqual(r.status_code, 413)
        self.assertFalse(RegistroAuditoria.objects.exists())

    def test_fallo_auditoria_no_entrega_archivo(self):
        with patch('ordenes.models.RegistroAuditoria.objects.create', side_effect=RuntimeError('simulado')):
            self.assertEqual(self.exportar('inventario').status_code, 500)

    def test_filtros_invalidos_rechazados_sin_archivo(self):
        for tipo, filtro in [('ventas', {'periods': 'foo'}), ('recibos', {'fecha_inicio': 'ayer'}), ('egresos', {'proveedores': '9' * 50}), ('inventario', {'fecha_inicio': '2026-10-10', 'fecha_fin': '2026-10-01'})]:
            self.assertEqual(self.exportar(tipo, filtro).status_code, 400)
        self.assertFalse(RegistroAuditoria.objects.exists())
