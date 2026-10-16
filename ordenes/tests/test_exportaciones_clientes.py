from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from django.db import connection
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook
from rest_framework.test import APIClient

from ordenes.models import Cliente, ObservacionCliente, RegistroAuditoria
from ordenes.tests.utils import LottusTestCase


class DatosClientesTests(LottusTestCase):
    def setUp(self):
        self.admin = self.crear_usuario('admin-datos', role='administrador')
        self.vendedor = self.crear_usuario('vendedor-datos')
        self.otro = self.crear_usuario('otro-datos')
        self.permisos('vendedor', ['VER_CLIENTES', 'VER_PROPIAS_VENTAS', 'EXPORTAR_CLIENTES', 'EDITAR_CLIENTE'])
        self.propio = self.cliente('001234', 'Cliente propio')
        self.venta(self.vendedor, cliente=self.propio)
        self.compartido = self.cliente('005678', 'Cliente compartido')
        self.venta(self.otro, cliente=self.compartido).vendedores_compartidos.add(self.vendedor)
        self.ajeno = self.cliente('009999', 'Cliente ajeno')
        self.venta(self.otro, cliente=self.ajeno)
        self.sin_ventas = self.cliente('008888', 'Sin ventas')
        self.api_vendedor = self.api(self.vendedor)

    def exportar(self, api=None, consulta=''):
        return (api or self.api_vendedor).post('/api/clientes/exportar/' + consulta, {}, format='json')

    def hoja(self, response):
        self.assertEqual(response.status_code, 200, getattr(response, 'data', ''))
        return load_workbook(BytesIO(response.content), data_only=False).active

    def test_exportacion_exige_ambos_permisos_y_no_admite_get(self):
        self.assertEqual(APIClient().post('/api/clientes/exportar/').status_code, 401)
        for codigos in (['VER_CLIENTES'], ['EXPORTAR_CLIENTES']):
            self.permisos('vendedor', codigos)
            self.assertEqual(self.exportar().status_code, 403)
        self.assertFalse(RegistroAuditoria.objects.exists())
        self.assertEqual(self.api(self.admin).get('/api/clientes/exportar/').status_code, 405)

    def test_descarga_solo_propios_compartidos_sin_duplicar_por_vendedores(self):
        self.propio.ventas.first().vendedores_compartidos.add(self.vendedor, self.otro)
        r = self.exportar()
        nombres = [fila[1].value for fila in list(self.hoja(r).rows)[1:]]
        self.assertCountEqual(nombres, ['Cliente propio', 'Cliente compartido'])
        self.assertEqual(r['Cache-Control'], 'no-store')
        self.assertEqual(r['X-Content-Type-Options'], 'nosniff')
        registro = RegistroAuditoria.objects.get(accion='clientes.exportar')
        self.assertEqual(registro.usuario, self.vendedor)
        self.assertEqual(registro.detalle['filas'], 2)
        self.assertEqual(registro.detalle['alcance'], 'propias_compartidas')

    def test_admin_exporta_global_y_preserva_fichas_sin_ventas(self):
        self.assertEqual(self.hoja(self.exportar(self.api(self.admin))).max_row, 5)

    def test_excel_texto_literal_con_ceros_telefonos_y_formulas(self):
        valores = {'nombre': '=HYPERLINK("https://example.invalid")', 'direccion': '\t@SUM(1,2)',
                   'barrio': '\ufeff=1+1', 'ciudad': '-1+2', 'telefono1': '+573012345678', 'telefono2': '0012345'}
        Cliente.objects.filter(pk=self.propio.pk).update(**valores)
        r = self.exportar(consulta='?cedula=001234')
        hoja = self.hoja(r)
        celdas = list(hoja.rows)[1]
        self.assertEqual([c.value for c in celdas], [str(self.propio.pk), valores['nombre'], '001234', None,
            valores['direccion'], valores['barrio'], valores['ciudad'], valores['telefono1'], valores['telefono2']])
        for celda in celdas:
            self.assertNotEqual(celda.data_type, 'f')
            self.assertIsNone(celda.hyperlink)
        with ZipFile(BytesIO(r.content)) as libro:
            self.assertNotIn(b'<f>', libro.read('xl/worksheets/sheet1.xml'))
        self.propio.refresh_from_db()
        self.assertEqual(self.propio.nombre, valores['nombre'])

    def test_filtro_se_aplica_completo_ignorando_paginacion(self):
        Cliente.objects.bulk_create([Cliente(nombre=f'Prueba lote {n}', cedula=f'lote-{n}', telefono1='0') for n in range(105)])
        r = self.exportar(self.api(self.admin), '?search=Prueba%20lote&page=3&page_size=1')
        self.assertEqual(self.hoja(r).max_row, 106)
        r = self.exportar(consulta=f'?query={self.propio.pk}&search=propio&cedula=001234')
        self.assertEqual(self.hoja(r).max_row, 2)

    def test_limite_rechaza_sin_truncar_y_sin_registrar_descarga_exitosa(self):
        with patch('ordenes.clientes_exportacion.LIMITE_EXPORTACION', 1):
            self.assertEqual(self.exportar().status_code, 413)
        self.assertFalse(RegistroAuditoria.objects.exists())

    def test_auditoria_no_copia_busqueda_documentos_ni_parametros_arbitrarios(self):
        self.exportar(consulta='?search=Cliente%20propio&cedula=001234&extra=secreto-prueba')
        detalle = RegistroAuditoria.objects.get().detalle
        texto = str(detalle)
        for dato in ('Cliente propio', '001234', 'secreto-prueba'):
            self.assertNotIn(dato, texto)
        self.assertEqual(set(detalle['filtros']), {'search', 'cedula'})
        self.assertEqual(len(detalle['filtros']['cedula']['huella']), 64)

    def test_error_de_auditoria_no_entrega_archivo_sin_trazabilidad(self):
        with patch('ordenes.clientes_exportacion.RegistroAuditoria.objects.create', side_effect=RuntimeError('Fallo simulado')):
            self.assertEqual(self.exportar().status_code, 500)
        self.assertFalse(RegistroAuditoria.objects.exists())

    def test_controles_invalidos_de_xml_se_rechazan_sin_modificar_el_original(self):
        Cliente.objects.filter(pk=self.propio.pk).update(nombre='Prueba\x00literal')
        self.assertEqual(self.exportar().status_code, 400)
        self.assertFalse(RegistroAuditoria.objects.exists())
        self.propio.refresh_from_db()
        self.assertEqual(self.propio.nombre, 'Prueba\x00literal')

    def test_filtros_validos_compartidos_y_numeros_fuera_de_rango(self):
        for ruta in ('/api/clientes/', '/api/clientes/exportar/'):
            metodo = self.api_vendedor.get if ruta.endswith('/clientes/') else self.api_vendedor.post
            self.assertEqual(metodo(ruta + '?search=' + '9' * 200).status_code, 200)
            self.assertEqual(metodo(ruta + '?search=' + '9' * 256).status_code, 400)

    def test_directorio_basico_global_se_conserva_y_explica_alcance_exportacion(self):
        r = self.api_vendedor.get('/api/clientes/?cedula=009999')
        self.assertEqual(r.data['results'][0]['id'], self.ajeno.pk)
        self.assertEqual(r.data['exportacion'], {'permitida': True, 'alcance': 'propias_compartidas'})
        self.assertEqual(self.api_vendedor.get(f'/api/clientes/{self.ajeno.pk}/').status_code, 200)
        self.assertEqual(self.hoja(self.exportar(consulta='?cedula=009999')).max_row, 1)

    def test_historial_no_revela_ventas_ajenas_del_mismo_cliente(self):
        ajena = self.venta(self.otro, cliente=self.propio)
        compartida = self.venta(self.otro, cliente=self.propio)
        compartida.vendedores_compartidos.add(self.vendedor)
        r = self.api_vendedor.get(f'/api/clientes/{self.propio.pk}/ventas-observaciones/')
        self.assertEqual(len(r.data['ventas']), 2)
        self.assertNotIn(ajena.pk, [v['id'] for v in r.data['ventas']])
        self.assertTrue(r.data['puede_agregar_observaciones'])
        self.assertTrue(r.data['historial_limitado'])
        global_r = self.api(self.admin).get(f'/api/clientes/{self.propio.pk}/ventas-observaciones/')
        self.assertEqual(len(global_r.data['ventas']), 3)
        self.assertFalse(global_r.data['historial_limitado'])

    def test_cliente_ajeno_solo_datos_basicos_sin_historial_ni_notas(self):
        obs = ObservacionCliente.objects.create(cliente=self.ajeno, autor=self.otro, texto='Nota privada ficticia')
        r = self.api_vendedor.get(f'/api/clientes/{self.ajeno.pk}/ventas-observaciones/')
        self.assertEqual(r.data['ventas'], [])
        self.assertEqual(r.data['observaciones_cliente'], [])
        self.assertTrue(r.data['historial_limitado'])
        self.assertFalse(r.data['puede_agregar_observaciones'])
        self.assertEqual(self.api_vendedor.post(f'/api/clientes/{self.ajeno.pk}/observaciones/anadir/', {'texto': 'Cambio'}).status_code, 404)
        for metodo in ('put', 'delete'):
            r = getattr(self.api_vendedor, metodo)(f'/api/observaciones-cliente/{obs.pk}/', {'texto': 'Cambio'})
            self.assertEqual(r.status_code, 404)
        obs.refresh_from_db()
        self.assertEqual(obs.texto, 'Nota privada ficticia')

    def test_permiso_clientes_no_revela_importes_de_ventas(self):
        self.permisos('vendedor', ['VER_CLIENTES'])
        r = self.api_vendedor.get(f'/api/clientes/{self.propio.pk}/ventas-observaciones/')
        self.assertEqual(r.data['ventas'], [])
        self.assertTrue(r.data['historial_limitado'])

    def test_retirar_alcance_quita_notas_incluso_a_su_autor(self):
        obs = ObservacionCliente.objects.create(cliente=self.compartido, autor=self.vendedor, texto='Mi nota')
        self.compartido.ventas.first().vendedores_compartidos.clear()
        self.assertEqual(self.api_vendedor.delete(f'/api/observaciones-cliente/{obs.pk}/').status_code, 404)
        self.assertTrue(ObservacionCliente.objects.filter(pk=obs.pk).exists())

    def test_historial_no_genera_consultas_por_venta(self):
        url = f'/api/clientes/{self.propio.pk}/ventas-observaciones/'
        with CaptureQueriesContext(connection) as pocas:
            self.api_vendedor.get(url)
        for _ in range(10):
            self.venta(self.vendedor, cliente=self.propio)
        with CaptureQueriesContext(connection) as muchas:
            self.api_vendedor.get(url)
        self.assertEqual(len(pocas), len(muchas))
        self.assertLessEqual(len(muchas), 15)
