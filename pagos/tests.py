from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from django.utils import timezone
from django.core.management import call_command
from openpyxl import load_workbook

from ordenes.models import ComprobanteEgreso, Caja
from ordenes.tests.utils import LottusTestCase
from .exportadores import ENCABEZADOS
from .models import CuentaBancaria, EmpleadoPago, LotePago, DetallePago, ProgramacionNomina
from .servicios import preparar_programada


class PagosFixture(LottusTestCase):
    def setUp(self):
        self.admin = self.crear_usuario('pagos_admin', role='administrador')
        self.api_admin = self.api(self.admin)
        self.prov = self.proveedor('Proveedor prueba bancaria')
        self.otro = self.proveedor('Otro proveedor prueba bancaria')
        self.origen = self.cuenta(empresa=True, numero='000012345678')
        self.destino = self.cuenta(proveedor=self.prov.id)

    def cuenta(self, **extra):
        d = {'banco': '51', 'tipo': 'CA', 'numero': '000456789012', 'tipo_documento': '01', 'documento': '0001234567',
             'nombres': 'José María', 'apellidos': 'Peña López', 'principal': True}
        d.update(extra)
        r = self.api_admin.post('/api/pagos-bancarios/cuentas/', d, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        return CuentaBancaria.objects.get(pk=r.data['id'])

    def ce(self, **extra):
        d = {'proveedor': self.prov.id, 'recibido_por': 'Persona de prueba', 'medio_pago': 'Transferencia',
             'valor': '1500000', 'cuenta_destino': self.destino.id}
        d.update(extra)
        r = self.api_admin.post('/api/comprobantes-egreso/crear/', d, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        return ComprobanteEgreso.objects.get(pk=r.data['id'])

    def lote(self, ce=None):
        ce = ce or self.ce()
        r = self.api_admin.post('/api/pagos-bancarios/lotes/', {'tipo': 'proveedores', 'fecha_pago': str(timezone.localdate()), 'comprobantes': [ce.id]}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        d = r.data
        r = self.api_admin.patch(f"/api/pagos-bancarios/lotes/{d['id']}/", {'revision': d['revision'], 'cuenta_origen': self.origen.id}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        return r.data

    def exportar(self, d):
        return self.api_admin.post(f"/api/pagos-bancarios/lotes/{d['id']}/archivo/", {'revision': d['revision'], 'huella_revision': d.get('huella_revision')}, format='json')


class PagosTests(PagosFixture):
    def test_excel_exactitud_ceros_nombres_columnas_hoja_sin_formulas(self):
        d = self.lote()
        r = self.exportar(d)
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(BytesIO(r.content), data_only=False)
        self.assertEqual(wb.sheetnames, ['Pagos'])
        ws = wb.active
        self.assertEqual([c.value for c in ws[1]], ENCABEZADOS)
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws.max_column, 9)
        self.assertEqual(ws['B2'].value, '0001234567')
        self.assertEqual(ws['G2'].value, '000456789012')
        self.assertEqual(ws['G2'].number_format, '@')
        self.assertEqual(ws['C2'].value, 'JOSE MARIA')
        self.assertEqual(ws['D2'].value, 'PENA LOPEZ')
        self.assertEqual(ws['E2'].value, 51)
        self.assertEqual(ws['F2'].value, 'CA')
        self.assertEqual(ws['H2'].value, 1500000)
        self.assertTrue(all(c.data_type != 'f' for row in ws for c in row))
        ce = ComprobanteEgreso.objects.get(pk=d['pagos'][0]['comprobante'])
        self.assertEqual(ce.estado, 'Por Confirmar Pago')
        self.assertEqual(Caja.objects.count(), 0)

    def test_multiples_proveedores_mismo_archivo(self):
        otra = self.cuenta(proveedor=self.otro.id, numero='12345', banco='7')
        a, b = self.ce(), self.ce(proveedor=self.otro.id, cuenta_destino=otra.id, valor='2100000')
        r = self.api_admin.post('/api/pagos-bancarios/lotes/', {'tipo': 'proveedores', 'fecha_pago': str(timezone.localdate()), 'comprobantes': [a.id,b.id]}, format='json')
        self.assertEqual(r.status_code, 201)
        d = self.api_admin.patch(f"/api/pagos-bancarios/lotes/{r.data['id']}/", {'revision': r.data['revision'], 'cuenta_origen': self.origen.id}, format='json').data
        wb = load_workbook(BytesIO(self.exportar(d).content))
        self.assertEqual(wb.active.max_row, 3)
        self.assertEqual(sum(wb.active.cell(i,8).value for i in (2,3)), 3600000)

    def test_cuenta_de_otro_proveedor_rechaza_ce_sin_escrituras(self):
        r = self.api_admin.post('/api/comprobantes-egreso/crear/', {'proveedor': self.otro.id, 'recibido_por': 'Prueba', 'medio_pago': 'Transferencia', 'valor': '10', 'cuenta_destino': self.destino.id}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertFalse(ComprobanteEgreso.objects.exists())

    def test_datos_incorrectos_de_cuenta_rechazados_sin_corregir_digitos(self):
        base = {'proveedor': self.prov.id, 'banco': '51', 'tipo': 'CA', 'tipo_documento': '03', 'documento': '9001234567', 'nombres': 'EMPRESA', 'apellidos': 'SAS'}
        for numero in ('', '123-456', '123.456', '12345678901234567', '1e6'):
            r = self.api_admin.post('/api/pagos-bancarios/cuentas/', {**base, 'numero': numero}, format='json')
            self.assertEqual(r.status_code, 400, r.data)

    def test_solo_una_cuenta_principal_y_archivar_conserva_historial(self):
        nueva = self.cuenta(proveedor=self.prov.id, numero='987654321')
        self.destino.refresh_from_db()
        self.assertFalse(self.destino.principal)
        self.assertTrue(nueva.principal)
        d = self.lote()
        self.assertEqual(self.exportar(d).status_code, 200)
        datos = {'proveedor': self.prov.id, 'revision': self.destino.revision, 'activa': False}
        self.assertEqual(self.api_admin.patch(f'/api/pagos-bancarios/cuentas/{self.destino.id}/', datos, format='json').status_code, 200)
        self.assertEqual(load_workbook(BytesIO(self.exportar(d).content)).active['G2'].value, '000456789012')

    def test_no_reserva_egreso_dos_veces(self):
        ce = self.ce()
        self.lote(ce)
        r = self.api_admin.post('/api/pagos-bancarios/lotes/', {'tipo': 'proveedores', 'fecha_pago': str(timezone.localdate()), 'comprobantes': [ce.id]}, format='json')
        self.assertEqual(r.status_code, 409)
        self.assertEqual(LotePago.objects.count(), 1)

    def test_reintento_crear_lote_misma_clave(self):
        ce = self.ce()
        d = {'tipo': 'proveedores', 'fecha_pago': str(timezone.localdate()), 'comprobantes': [ce.id]}
        a = self.api_admin.post('/api/pagos-bancarios/lotes/', d, format='json', HTTP_IDEMPOTENCY_KEY='prueba-lote')
        b = self.api_admin.post('/api/pagos-bancarios/lotes/', d, format='json', HTTP_IDEMPOTENCY_KEY='prueba-lote')
        self.assertEqual(a.status_code, 201)
        self.assertEqual(a.data['id'], b.data['id'])
        self.assertEqual(LotePago.objects.count(), 1)

    def test_revision_concurrente_y_atomicidad_de_edicion(self):
        d = self.lote()
        url = f"/api/pagos-bancarios/lotes/{d['id']}/"
        r = self.api_admin.patch(url, {'revision': 0, 'nombre': 'No guardar'}, format='json')
        self.assertEqual(r.status_code, 409)
        r = self.api_admin.patch(url, {'revision': d['revision'], 'nombre': 'No guardar', 'pagos': [{'id': d['pagos'][0]['id'], 'valor': '999'}]}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertNotEqual(LotePago.objects.get(pk=d['id']).nombre, 'No guardar')

    def test_falla_exportacion_revierte_estado_y_snapshots(self):
        d = self.lote()
        with patch('pagos.servicios.exportar', side_effect=RuntimeError('Exportador de prueba')):
            self.assertEqual(self.exportar(d).status_code, 500)
        lote = LotePago.objects.get(pk=d['id'])
        self.assertEqual(lote.estado, 'borrador')
        self.assertFalse(lote.archivo)
        self.assertEqual(lote.pagos.get().cuenta_snapshot, {})

    def test_archivo_inmutable_tras_editar_maestro_y_reintentar(self):
        d = self.lote()
        a = self.exportar(d)
        self.destino.numero = '999999'
        self.destino.save()
        b = self.exportar(d)
        self.assertEqual(a.content, b.content)
        self.assertEqual(ComprobanteEgreso.objects.count(), 1)

    def test_ce_exportado_no_se_edita_pero_se_confirma(self):
        ce = self.ce()
        d = self.lote(ce)
        self.exportar(d)
        r = self.api_admin.put(f'/api/comprobantes-egreso/{ce.id}/editar/', {}, format='json')
        self.assertEqual(r.status_code, 409)
        r = self.api_admin.patch(f'/api/comprobantes-egreso/{ce.id}/confirmar/', {}, format='json')
        self.assertEqual(r.status_code, 200)
        ce.refresh_from_db()
        self.assertEqual(ce.estado, 'Pagado')

    def test_cuenta_inactiva_o_ce_cambiado_impide_descargar(self):
        d = self.lote()
        self.destino.activa = False
        self.destino.save()
        r = self.exportar(d)
        self.assertEqual(r.status_code, 400)
        self.destino.activa = True
        self.destino.save()
        ComprobanteEgreso.objects.update(valor='1900000')
        self.assertEqual(self.exportar(d).status_code, 400)
        self.assertEqual(LotePago.objects.get(pk=d['id']).estado, 'borrador')

    def test_cambio_de_cuenta_en_ce_se_detecta_en_lote(self):
        d = self.lote()
        ComprobanteEgreso.objects.update(cuenta_destino=None)
        self.assertEqual(self.exportar(d).status_code, 400)

    def test_bancolombia_destino_admitido_origen_bloqueado(self):
        origen = self.cuenta(empresa=True, banco='7', numero='98765654')
        d = self.lote()
        r = self.api_admin.patch(f"/api/pagos-bancarios/lotes/{d['id']}/", {'revision': d['revision'], 'cuenta_origen': origen.id}, format='json')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.exportar(r.data).status_code, 400)

    def test_anulacion_exige_archivo_no_procesado_y_libera_egresos(self):
        ce = self.ce()
        d = self.lote(ce)
        self.exportar(d)
        lote = LotePago.objects.get(pk=d['id'])
        url = f'/api/pagos-bancarios/lotes/{lote.id}/anular/'
        datos = {'revision': lote.revision, 'motivo': 'Archivo de prueba no cargado'}
        self.assertEqual(self.api_admin.post(url, datos, format='json').status_code, 400)
        self.assertEqual(self.api_admin.post(url, {**datos, 'archivo_no_procesado': True}, format='json').status_code, 200)
        self.assertEqual(self.exportar(d).status_code, 409)
        self.assertEqual(self.api_admin.get('/api/pagos-bancarios/pendientes/').data[0]['id'], ce.id)

    def test_pago_confirmado_no_se_anula(self):
        d = self.lote()
        self.exportar(d)
        ComprobanteEgreso.objects.update(estado='Pagado')
        lote = LotePago.objects.get(pk=d['id'])
        r = self.api_admin.post(f'/api/pagos-bancarios/lotes/{lote.id}/anular/', {'revision': lote.revision, 'motivo': 'Prueba', 'archivo_no_procesado': True}, format='json')
        self.assertEqual(r.status_code, 409)

    def test_no_exporta_si_cambiaron_datos_de_cuenta_despues_de_revisar(self):
        d = self.lote()
        self.destino.numero = '111222333'
        self.destino.save(update_fields=['numero'])
        self.assertEqual(self.exportar(d).status_code, 409)
        self.assertFalse(LotePago.objects.get(pk=d['id']).archivo)

    def test_confirmacion_parcial_no_paga_los_demas_y_reintentar_no_duplica(self):
        a,b = self.ce(), self.ce(valor='2300000')
        r = self.api_admin.post('/api/pagos-bancarios/lotes/', {'tipo':'proveedores','fecha_pago':str(timezone.localdate()),'comprobantes':[a.id,b.id]},format='json')
        d = self.api_admin.patch(f"/api/pagos-bancarios/lotes/{r.data['id']}/", {'revision':r.data['revision'],'cuenta_origen':self.origen.id},format='json').data
        self.assertEqual(self.exportar(d).status_code,200)
        url = f"/api/pagos-bancarios/lotes/{d['id']}/confirmar/"
        datos = {'comprobantes':[a.id], 'referencia':'OPERACION PRUEBA 123'}
        self.assertEqual(self.api_admin.post(url,datos,format='json').status_code,200)
        a.refresh_from_db(); b.refresh_from_db()
        self.assertEqual(a.estado,'Pagado'); self.assertEqual(b.estado,'Por Confirmar Pago')
        self.assertEqual(self.api_admin.post(url,datos,format='json').data['confirmados'],1)
        fuera = self.ce()
        self.assertEqual(self.api_admin.post(url,{**datos,'comprobantes':[fuera.id]},format='json').status_code,400)

    def test_aprobar_pago_requiere_permiso_adicional(self):
        d = self.lote()
        self.assertEqual(self.exportar(d).status_code,200)
        self.permisos('auxiliar',['VER_COMPROBANTES_EGRESO','VER_TODOS_EGRESOS','CREAR_COMPROBANTE_EGRESO','EXPORTAR_EGRESOS'])
        c = self.api(self.crear_usuario('aux_sin_aprobacion',role='auxiliar'))
        r=c.post(f"/api/pagos-bancarios/lotes/{d['id']}/confirmar/",{'comprobantes':[d['pagos'][0]['comprobante']], 'referencia':'PRUEBA'},format='json')
        self.assertEqual(r.status_code,403)
        self.assertEqual(ComprobanteEgreso.objects.get().estado,'Por Confirmar Pago')

    def test_reintento_parcial_exige_rechazo_no_duplica_ce_y_omite_confirmados(self):
        a,b = self.ce(),self.ce(valor='2300000')
        r=self.api_admin.post('/api/pagos-bancarios/lotes/',{'tipo':'proveedores','fecha_pago':str(timezone.localdate()),'comprobantes':[a.id,b.id]},format='json')
        d=self.api_admin.patch(f"/api/pagos-bancarios/lotes/{r.data['id']}/",{'revision':r.data['revision'],'cuenta_origen':self.origen.id},format='json').data
        self.assertEqual(self.exportar(d).status_code,200)
        self.api_admin.post(f"/api/pagos-bancarios/lotes/{d['id']}/confirmar/",{'comprobantes':[a.id],'referencia':'EXITO PARCIAL'},format='json')
        lote=LotePago.objects.get(pk=d['id'])
        url=f'/api/pagos-bancarios/lotes/{lote.id}/reintentar/'
        datos={'revision':lote.revision,'comprobantes':[b.id],'referencia':'RECHAZO BANCARIO'}
        self.assertEqual(self.api_admin.post(url,datos,format='json').status_code,400)
        self.assertEqual(self.api_admin.post(url,{**datos,'comprobantes':[a.id],'rechazo_confirmado':True},format='json').status_code,409)
        nuevo=self.api_admin.post(url,{**datos,'rechazo_confirmado':True},format='json')
        self.assertEqual(nuevo.status_code,200,nuevo.data)
        self.assertEqual(len(nuevo.data['pagos']),1)
        self.assertEqual(nuevo.data['pagos'][0]['comprobante'],b.id)
        exportado=self.exportar(nuevo.data)
        self.assertEqual(exportado.status_code,200)
        ws=load_workbook(BytesIO(exportado.content)).active
        self.assertEqual(ws.max_row,2)
        self.assertEqual(ws['I2'].value,str(b.id))
        self.assertEqual(ComprobanteEgreso.objects.count(),2)
        self.assertEqual(self.api_admin.post(url,{**datos,'rechazo_confirmado':True},format='json').status_code,409)

    def test_cuentas_y_nomina_no_expuestas_a_vendedor(self):
        vendedor = self.api(self.crear_usuario('vendedor_bancos', role='vendedor'))
        for url in ('/api/pagos-bancarios/empleados/', '/api/pagos-bancarios/cuentas/?empresa=1', '/api/pagos-bancarios/lotes/?tipo=nomina', '/api/pagos-bancarios/pendientes/'):
            self.assertEqual(vendedor.get(url).status_code, 403)

    def test_reintento_idempotente_no_elude_un_permiso_revocado(self):
        u=self.crear_usuario('aux_cuenta',role='auxiliar')
        self.permisos('auxiliar',['CREAR_COMPROBANTE_EGRESO'])
        d={'proveedor':self.prov.id,'banco':'51','tipo':'CA','numero':'000777666555','tipo_documento':'01','documento':'123456789','nombres':'Persona','apellidos':'Ficticia'}
        c=self.api(u)
        self.assertEqual(c.post('/api/pagos-bancarios/cuentas/',d,format='json',HTTP_IDEMPOTENCY_KEY='cuenta-permiso').status_code,201)
        self.permisos('auxiliar',[])
        self.assertEqual(c.post('/api/pagos-bancarios/cuentas/',d,format='json',HTTP_IDEMPOTENCY_KEY='cuenta-permiso').status_code,403)


class NominaBancariaTests(PagosFixture):
    def nomina(self):
        e = EmpleadoPago.objects.create(nombre='Empleado ficticio', pago_mensual='1800000')
        self.cuenta(empleado=e.id)
        r = self.api_admin.post('/api/pagos-bancarios/lotes/', {'tipo': 'nomina', 'periodo': timezone.localdate().strftime('%Y-%m')}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        d = r.data
        r = self.api_admin.patch(f"/api/pagos-bancarios/lotes/{d['id']}/", {'revision': d['revision'], 'cuenta_origen': self.origen.id, 'fecha_pago': str(timezone.localdate())}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        return e, r.data

    def test_nomina_ajustable_crea_ce_solo_al_exportar_y_sin_duplicados(self):
        e, d = self.nomina()
        self.assertFalse(ComprobanteEgreso.objects.exists())
        d = self.api_admin.patch(f"/api/pagos-bancarios/lotes/{d['id']}/", {'revision': d['revision'], 'pagos': [{'id': d['pagos'][0]['id'], 'valor': '1950000'}]}, format='json').data
        r = self.exportar(d)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(r.content)).active['H2'].value, 1950000)
        e.refresh_from_db()
        self.assertEqual(e.pago_mensual, Decimal('1800000'))
        self.assertEqual(ComprobanteEgreso.objects.get().estado, 'Por Confirmar Pago')
        self.assertEqual(self.exportar(d).content, r.content)
        self.assertEqual(ComprobanteEgreso.objects.count(), 1)
        self.assertFalse(Caja.objects.exists())

    def test_ceros_fila_excluida_y_decimales(self):
        e,d = self.nomina()
        url = f"/api/pagos-bancarios/lotes/{d['id']}/"
        for valor in ('', '0.5', '-1', 'nan', '1e200', '1.000.000'):
            r = self.api_admin.patch(url, {'revision': d['revision'], 'pagos': [{'id': d['pagos'][0]['id'], 'valor': valor}]}, format='json')
            self.assertEqual(r.status_code, 400, r.data)
        r = self.api_admin.patch(url, {'revision': d['revision'], 'pagos': [{'id': d['pagos'][0]['id'], 'valor': '0'}]}, format='json')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.exportar(r.data).status_code, 400)

    def test_falla_exportador_nomina_no_deja_egresos(self):
        e,d = self.nomina()
        with patch('pagos.servicios.exportar', side_effect=RuntimeError('Prueba')):
            self.assertEqual(self.exportar(d).status_code, 500)
        self.assertFalse(ComprobanteEgreso.objects.exists())
        self.assertEqual(LotePago.objects.get(pk=d['id']).estado, 'borrador')

    def test_programacion_dia_31_febrero_idempotente_conserva_ajustes(self):
        e = EmpleadoPago.objects.create(nombre='Persona programada', pago_mensual='1000000')
        self.cuenta(empleado=e.id)
        cfg = ProgramacionNomina.objects.get(pk=1)
        cfg.activa, cfg.dia, cfg.cuenta_origen, cfg.actualizado_por, cfg.habilitada_desde = True, 31, self.origen, self.admin, date(2026, 1, 1)
        cfg.save()
        self.assertIsNone(preparar_programada(date(2026,2,27)))
        d = preparar_programada(date(2026,2,28))
        self.assertEqual(d.fecha_pago, date(2026,2,28))
        DetallePago.objects.filter(lote=d).update(valor='1200000')
        self.assertEqual(preparar_programada(date(2026,2,28)).id, d.id)
        self.assertEqual(d.pagos.get().valor, Decimal('1200000'))
        self.assertEqual(LotePago.objects.count(), 1)
        self.assertFalse(ComprobanteEgreso.objects.exists())

    def test_dos_cuentas_sin_principal_exigen_seleccion(self):
        e,d = self.nomina()
        self.cuenta(empleado=e.id, numero='54321234', principal=False)
        CuentaBancaria.objects.filter(empleado=e).update(principal=False)
        otro_mes = (timezone.localdate().replace(day=1) + timedelta(days=32)).strftime('%Y-%m')
        r = self.api_admin.post('/api/pagos-bancarios/lotes/', {'tipo': 'nomina', 'periodo': otro_mes}, format='json')
        self.assertEqual(r.status_code, 201)
        self.assertIsNone(r.data['pagos'][0]['cuenta'])

    def test_programacion_desactivada_y_validacion_dias(self):
        self.assertIsNone(preparar_programada(date(2026,10,25)))
        for dia in (0,32,'abc'):
            r = self.api_admin.put('/api/pagos-bancarios/programacion/', {'revision': 0, 'dia': dia}, format='json')
            self.assertEqual(r.status_code, 400)
        r = self.api_admin.put('/api/pagos-bancarios/programacion/', {'revision': 0, 'activa': True, 'dia': 25}, format='json')
        self.assertEqual(r.status_code, 400)

    def test_corregir_nomina_reutiliza_ce_y_conserva_archivos_anteriores(self):
        e,d = self.nomina()
        original = self.exportar(d)
        self.assertEqual(original.status_code,200)
        detalle = self.api_admin.get(f"/api/pagos-bancarios/lotes/{d['id']}/").data
        anulado = self.api_admin.post(f"/api/pagos-bancarios/lotes/{d['id']}/anular/", {'revision':detalle['revision'],'motivo':'Corregir antes de cargar','archivo_no_procesado':True},format='json')
        self.assertEqual(anulado.status_code,200)
        corregir = self.api_admin.post(f"/api/pagos-bancarios/lotes/{d['id']}/corregir/", {'revision':anulado.data['revision']},format='json')
        self.assertEqual(corregir.status_code,200)
        nuevo = self.api_admin.patch(f"/api/pagos-bancarios/lotes/{d['id']}/", {'revision':corregir.data['revision'],'pagos':[{'id':d['pagos'][0]['id'],'valor':'1990000'}]},format='json').data
        segundo = self.exportar(nuevo)
        self.assertEqual(segundo.status_code,200)
        self.assertEqual(ComprobanteEgreso.objects.count(),1)
        self.assertEqual(ComprobanteEgreso.objects.get().valor,Decimal('1990000'))
        self.assertEqual(LotePago.objects.get(pk=d['id']).archivos.count(),2)
        self.assertNotEqual(segundo.content,original.content)

    def test_motor_continuo_se_recupera_tras_error_temporal(self):
        from django.db import OperationalError
        with patch('pagos.management.commands.preparar_pagos_nomina.preparar_programada',side_effect=[OperationalError('Prueba temporal'),None]) as preparar:
            with patch('pagos.management.commands.preparar_pagos_nomina.time.sleep',side_effect=[None,KeyboardInterrupt]):
                with self.assertLogs('pagos.management.commands.preparar_pagos_nomina',level='ERROR'):
                    with self.assertRaises(KeyboardInterrupt):
                        call_command('preparar_pagos_nomina',continuo=True)
            self.assertEqual(preparar.call_count,2)


class PreparacionDesdeCETests(PagosFixture):
    def revisar(self, ces):
        r = self.api_admin.post('/api/pagos-bancarios/revisar-ce/', {'comprobantes': [c.id for c in ces]}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        return r.data

    def test_davivienda_nombre_y_directorio_asociacion_explicita(self):
        self.assertEqual(self.api_admin.get('/api/pagos-bancarios/catalogo/').data['bancos']['51'], 'Davivienda')
        data = self.api_admin.get('/api/pagos-bancarios/beneficiarios/').data
        p = next(b for b in data if b['titular'] == {'proveedor': self.prov.id})
        self.assertEqual(p['cuentas'], 1)
        self.assertEqual(p['bancos'], ['Davivienda'])

    def test_cuentas_telas_seleccion_desde_ce_y_excel(self):
        from ordenes.models import ProveedorTela
        tela = ProveedorTela.objects.create(nombre_empresa='Telas QA')
        cuenta = self.cuenta(proveedor_tela=tela.id, numero='000112244')
        ce = self.ce(proveedor='', proveedor_tela=tela.id, cuenta_destino=cuenta.id)
        data = self.revisar([ce])[0]
        self.assertEqual(data['proveedor_nombre'], 'Telas QA')
        self.assertEqual(data['pago_bancario']['titular'], {'proveedor_tela': tela.id})
        r = self.api_admin.post('/api/pagos-bancarios/lotes/', {'tipo': 'proveedores', 'fecha_pago': str(timezone.localdate()), 'cuenta_origen': self.origen.id, 'comprobantes': [ce.id], 'pagos': [{'comprobante': ce.id, 'cuenta': cuenta.id, 'huella': data['pago_bancario']['huella']}]}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(self.exportar(r.data).status_code, 200)
        self.assertEqual(ComprobanteEgreso.objects.get(pk=ce.pk).estado, 'Por Confirmar Pago')

    def test_filtro_banco_usa_cuenta_explicita_o_principal_y_sin_cuenta(self):
        dav = self.ce(cuenta_destino=None)
        banco = self.cuenta(proveedor=self.otro.id, banco='7', numero='123123')
        otro = self.ce(proveedor=self.otro.id, cuenta_destino=banco.id)
        for filtro, ids in [('51', [dav.id]), ('7', [otro.id]), ('sin_cuenta', [])]:
            r = self.api_admin.get('/api/comprobantes-egreso/', {'banco_destino': filtro, 'solo_pagables': '1'})
            self.assertEqual(r.status_code, 200, r.data)
            self.assertEqual([c['id'] for c in r.data['results']], ids)
        self.destino.principal = False; self.destino.save()
        self.cuenta(proveedor=self.prov.id, numero='112233', principal=False)
        r = self.api_admin.get('/api/comprobantes-egreso/', {'banco_destino': 'sin_cuenta', 'solo_pagables': '1'})
        self.assertEqual([c['id'] for c in r.data['results']], [dav.id])

    def test_asocia_nombre_libre_a_telas_sin_inferir_por_nombre(self):
        from ordenes.models import ProveedorTela
        tela = ProveedorTela.objects.create(nombre_empresa='Proveedor de tela')
        cuenta = self.cuenta(proveedor_tela=tela.id, numero='000345')
        ce = self.ce(proveedor='', proveedor_otro_nombre='Nombre anterior', cuenta_destino=None)
        d = self.revisar([ce])[0]
        self.assertIsNone(d['pago_bancario']['titular'])
        r = self.api_admin.patch(f'/api/pagos-bancarios/comprobantes/{ce.id}/beneficiario/', {'proveedor_tela': tela.id, 'cuenta': cuenta.id, 'huella': d['pago_bancario']['huella']}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        ce.refresh_from_db()
        self.assertEqual(ce.proveedor_tela_id, tela.id)
        self.assertEqual(ce.cuenta_destino_id, cuenta.id)
        self.assertEqual(ce.valor, Decimal('1500000'))

    def test_revision_obsoleta_rechaza_todo_sin_reservar_ni_cambiar_cuentas(self):
        ce = self.ce(cuenta_destino=None)
        d = self.revisar([ce])[0]
        ce.valor = Decimal('1600000'); ce.save()
        r = self.api_admin.post('/api/pagos-bancarios/lotes/', {'tipo': 'proveedores', 'fecha_pago': str(timezone.localdate()), 'cuenta_origen': self.origen.id, 'comprobantes': [ce.id], 'pagos': [{'comprobante': ce.id, 'cuenta': self.destino.id, 'huella': d['pago_bancario']['huella']}]}, format='json')
        self.assertEqual(r.status_code, 409, r.data)
        self.assertFalse(LotePago.objects.exists())
        ce.refresh_from_db(); self.assertIsNone(ce.cuenta_destino_id)

    def test_cuenta_ajena_rechaza_preparacion_y_hace_rollback(self):
        ce = self.ce(cuenta_destino=None)
        otra = self.cuenta(proveedor=self.otro.id, numero='556677')
        d = self.revisar([ce])[0]
        r = self.api_admin.post('/api/pagos-bancarios/lotes/', {'tipo': 'proveedores', 'fecha_pago': str(timezone.localdate()), 'comprobantes': [ce.id], 'pagos': [{'comprobante': ce.id, 'cuenta': otra.id, 'huella': d['pago_bancario']['huella']}]}, format='json')
        self.assertEqual(r.status_code, 400, r.data)
        self.assertFalse(LotePago.objects.exists())

    def test_preparado_no_se_selecciona_otra_vez_y_cuenta_se_guarda(self):
        ce = self.ce(cuenta_destino=None)
        d = self.revisar([ce])[0]
        r = self.api_admin.post('/api/pagos-bancarios/lotes/', {'tipo': 'proveedores', 'fecha_pago': str(timezone.localdate()), 'cuenta_origen': self.origen.id, 'comprobantes': [ce.id], 'pagos': [{'comprobante': ce.id, 'cuenta': self.destino.id, 'huella': d['pago_bancario']['huella']}]}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        ce.refresh_from_db(); self.assertEqual(ce.cuenta_destino_id, self.destino.id)
        info = self.revisar([ce])[0]['pago_bancario']
        self.assertFalse(info['disponible']); self.assertEqual(info['archivo'], r.data['id'])
        self.assertEqual(self.api_admin.get('/api/comprobantes-egreso/', {'solo_pagables': '1'}).data['count'], 0)

    def test_eliminar_cuenta_conserva_historico_y_quita_de_nuevos_pagos(self):
        d = self.lote()
        self.assertEqual(self.exportar(d).status_code, 200)
        cuenta = self.api_admin.get('/api/pagos-bancarios/cuentas/', {'proveedor': self.prov.id}).data[0]
        r = self.api_admin.delete(f'/api/pagos-bancarios/cuentas/{self.destino.id}/', cuenta, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.destino.refresh_from_db(); self.assertFalse(self.destino.activa)
        self.assertEqual(self.exportar(d).status_code, 200)
        self.assertEqual(ComprobanteEgreso.objects.count(), 1)

    def test_nomina_vincula_usuario_y_cuenta_sin_permitir_cambiar_persona(self):
        persona = self.crear_usuario('persona_nomina', role='auxiliar')
        r = self.api_admin.post('/api/pagos-bancarios/empleados/', {'nombre': 'Persona nómina', 'usuario': persona.id, 'pago_mensual': '2000000'}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        e = r.data
        c = self.cuenta(empleado=e['id'], numero='001122')
        listado = self.api_admin.get('/api/pagos-bancarios/beneficiarios/').data
        b = next(x for x in listado if x['titular'] == {'empleado': e['id']})
        self.assertEqual(b['usuario'], persona.id)
        self.assertEqual(b['cuentas'], 1)
        self.assertEqual(CuentaBancaria.objects.get(pk=c.pk).empleado.usuario_id, persona.id)
        r = self.api_admin.patch(f"/api/pagos-bancarios/empleados/{e['id']}/", {'usuario': self.admin.id}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertEqual(EmpleadoPago.objects.get(pk=e['id']).usuario_id, persona.id)
        r = self.api_admin.post('/api/pagos-bancarios/empleados/', {'nombre': 'Sin asociación', 'pago_mensual': '100'}, format='json')
        self.assertEqual(r.status_code, 400)

    def test_consulta_ce_no_hace_una_consulta_bancaria_por_cada_fila(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        self.ce()
        with CaptureQueriesContext(connection) as pocas:
            r = self.api_admin.get('/api/comprobantes-egreso/')
            self.assertEqual(r.status_code, 200)
        for i in range(10):
            p = self.proveedor('Proveedor consultas ' + str(i))
            c = self.cuenta(proveedor=p.id, numero=str(100000+i))
            self.ce(proveedor=p.id, cuenta_destino=c.id)
        with CaptureQueriesContext(connection) as muchas:
            r = self.api_admin.get('/api/comprobantes-egreso/')
            self.assertEqual(r.status_code, 200)
            self.assertEqual(r.data['count'], 11)
        self.assertLessEqual(len(muchas), len(pocas) + 2)
